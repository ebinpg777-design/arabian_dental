# Part of ebshel_einvoice_report. See LICENSE file for full copyright and licensing details.
import io

from odoo import Command
from odoo.exceptions import UserError
from odoo.tests import tagged

from odoo.addons.l10n_in.tests.common import L10nInTestInvoicingCommon

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


@tagged('post_install_l10n', 'post_install', '-at_install')
class TestEinvoiceReport(L10nInTestInvoicingCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.partner_a.write({'l10n_in_gst_treatment': 'regular', 'phone': '+91 98765 43210',
                             'email': 'clinic@example.com'})
        cls.partner_b.write({'l10n_in_gst_treatment': 'regular'})
        cls.consumer = cls.env['res.partner'].create({
            'name': 'Walk-in Patient', 'l10n_in_gst_treatment': 'consumer',
            'state_id': cls.default_company.state_id.id, 'country_id': cls.country_in.id})
        cls.product_a.write({'l10n_in_hsn_code': '9021.21.00'})
        cls.service = cls.env['product.product'].create({
            'name': 'Design Fee', 'type': 'service', 'lst_price': 100.0,
            'l10n_in_hsn_code': '998391', 'taxes_id': [Command.set(cls.sgst_sale_5.ids)]})
        cls.day = '2023-05-20'

        def invoice(partner, lines, move_type='out_invoice', post=True):
            move = cls.env['account.move'].create({
                'move_type': move_type, 'partner_id': partner.id, 'invoice_date': cls.day,
                'invoice_line_ids': [Command.create(dict({'quantity': 1.0}, **line)) for line in lines],
            })
            if post:
                move.action_post()
            return move

        # intra-state, two lines: a discounted product and a service
        cls.intra = invoice(cls.partner_a, [
            {'product_id': cls.product_a.id, 'quantity': 2.0, 'price_unit': 1000.0, 'discount': 10.0,
             'tax_ids': [Command.set(cls.sgst_sale_5.ids)]},
            {'product_id': cls.service.id, 'price_unit': 100.0, 'tax_ids': [Command.set(cls.sgst_sale_5.ids)]},
        ])
        # inter-state, IGST 18 %
        cls.inter = invoice(cls.partner_b, [
            {'product_id': cls.product_a.id, 'price_unit': 500.0, 'tax_ids': [Command.set(cls.igst_sale_18.ids)]},
        ])
        cls.to_consumer = invoice(cls.consumer, [
            {'product_id': cls.product_a.id, 'price_unit': 300.0, 'tax_ids': [Command.set(cls.sgst_sale_5.ids)]},
        ])
        cls.draft = invoice(cls.partner_a, [
            {'product_id': cls.product_a.id, 'price_unit': 300.0, 'tax_ids': [Command.set(cls.sgst_sale_5.ids)]},
        ], post=False)
        cls.refund = invoice(cls.partner_a, [
            {'product_id': cls.product_a.id, 'price_unit': 200.0, 'tax_ids': [Command.set(cls.sgst_sale_5.ids)]},
        ], move_type='out_refund')

    def _wizard(self, **vals):
        return self.env['einvoice.report.wizard'].create(dict({
            'date_from': self.day, 'date_to': self.day,
            'company_ids': [Command.set(self.default_company.ids)]}, **vals))

    def _rows(self, wizard):
        return {buyer['number']: (buyer, items) for buyer, items in wizard._documents()}

    def test_only_einvoiced_posted_invoices_and_credit_notes_on_request(self):
        rows = self._rows(self._wizard())
        self.assertEqual(set(rows), {self.intra.name, self.inter.name},
                         "registered customers only, posted only, no credit note by default")
        rows = self._rows(self._wizard(include_credit_notes=True))
        self.assertEqual(rows[self.refund.name][0]['doc_type'], 'CRN')
        self.assertIn('2 document(s)', self._wizard().summary)

    def test_the_buyer_row_reads_like_the_e_invoice(self):
        rows = self._rows(self._wizard())
        buyer, _items = rows[self.intra.name]
        self.assertEqual((buyer['supply_type'], buyer['igst_on_intra'], buyer['doc_type'], buyer['date']),
                         ('B2B', 'N', 'INV', '20/05/2023'))
        self.assertEqual((buyer['gstin'], buyer['name'], buyer['pos'], buyer['state_code'], buyer['pin']),
                         ('24ABCPM8965E1ZE', 'Partner Intra State', '24', '24', '360001'),
                         "GST state codes, not state names")
        self.assertEqual((buyer['phone'], buyer['email']), ('919876543210', 'clinic@example.com'))
        self.assertEqual(buyer['taxable'], 1900.0, "2 x 1000 less 10 %, plus 100")
        self.assertEqual((buyer['sgst'], buyer['cgst'], buyer['igst']), (47.5, 47.5, 0.0))
        self.assertEqual(buyer['discount'], 0.0, "line discounts are already out of the taxable value")
        for number, move in ((self.intra.name, self.intra), (self.inter.name, self.inter)):
            self.assertAlmostEqual(rows[number][0]['total'], move.amount_total, places=2,
                                   msg="taxable + taxes + round off is the invoice total")
        other = rows[self.inter.name][0]
        self.assertEqual((other['pos'], other['state_code'], other['igst'], other['sgst']),
                         ('27', '27', 90.0, 0.0))

    def test_the_item_rows_carry_the_line_split(self):
        _buyer, items = self._rows(self._wizard())[self.intra.name]
        product, service = items
        self.assertEqual((product['sl_no'], product['description'], product['hsn'], product['is_service']),
                         (1, self.product_a.name, '90212100', 'N'))
        self.assertEqual((product['quantity'], product['unit_price'], product['gross'], product['discount'],
                          product['taxable']), (2.0, 1000.0, 2000.0, 200.0, 1800.0))
        self.assertEqual((product['gst_rate'], product['sgst'], product['cgst'], product['total']),
                         (5.0, 45.0, 45.0, 1890.0))
        self.assertTrue(product['unit'], "a GST unit code (UQC), never blank")
        self.assertEqual((service['is_service'], service['hsn']), ('Y', '998391'),
                         "an SAC (99...) is a service")

    def test_the_workbook_has_both_sheets_and_totals_that_do_not_hide_a_column(self):
        if openpyxl is None:
            self.skipTest("openpyxl is not installed")
        data = self._wizard()._xlsx()
        workbook = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(workbook.sheetnames, ['Buyer Details', 'E-Invoice Items Sheet'])
        items = workbook['E-Invoice Items Sheet']
        headers = [cell.value for cell in items[7]][1:]
        self.assertEqual(headers[10], 'Gross Amount')
        total_row = [cell.value for cell in items[items.max_row]]
        self.assertEqual(total_row[1], 'Grand Total')
        self.assertEqual(total_row[1 + 10], 2000.0 + 100.0 + 500.0, "the gross total is shown, not merged over")
        buyer = workbook['Buyer Details']
        first = [cell.value for cell in buyer[8]][1:]
        self.assertEqual(first[11], '360001', "the PIN stays text")

    def test_one_upload_is_one_state(self):
        other = self.env['res.company'].create({'name': 'Other State Co', 'state_id': self.state_in_mh.id,
                                                'country_id': self.country_in.id})
        self.env.user.company_ids |= other
        wizard = self._wizard(company_ids=[Command.set((self.default_company | other).ids)])
        with self.assertRaises(UserError):
            wizard.action_download()
        with self.assertRaises(UserError):
            self._wizard(date_from='2023-05-21').action_download()
