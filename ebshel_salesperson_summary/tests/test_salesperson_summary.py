# Part of ebshel_salesperson_summary. See LICENSE file for full copyright and licensing details.
import io

from odoo import Command
from odoo.exceptions import UserError
from odoo.tests import tagged

from odoo.addons.account.tests.common import AccountTestInvoicingCommon

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


@tagged('post_install', '-at_install')
class TestSalespersonSummary(AccountTestInvoicingCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        group_user = cls.env.ref('base.group_user')
        cls.anita, cls.binu = cls.env['res.users'].create([
            {'name': 'Anita Sales', 'login': 'anita_sps', 'group_ids': [Command.set(group_user.ids)]},
            {'name': 'Binu Sales', 'login': 'binu_sps', 'group_ids': [Command.set(group_user.ids)]},
        ])
        Partner = cls.env['res.partner']
        cls.clinic_a = cls.partner_a
        cls.clinic_a.write({'user_id': cls.anita.id, 'city': 'Manjeri'})
        cls.clinic_b = cls.partner_b
        cls.clinic_b.user_id = cls.binu
        cls.walk_in = Partner.create({'name': 'Walk-in Clinic'})
        cls.settled = Partner.create({'name': 'Settled Clinic', 'user_id': cls.anita.id})
        cls.receivable = cls.company_data['default_account_receivable']
        cls.bank = cls.company_data['default_journal_bank']
        cls.misc = cls.company_data['default_journal_misc']
        no_tax = cls.env['account.tax']

        def invoice(partner, amount, day, move_type='out_invoice'):
            return cls._create_invoice(
                move_type=move_type, partner_id=partner, invoice_date=day, post=True,
                invoice_line_ids=[cls._prepare_invoice_line(price_unit=amount, tax_ids=no_tax)])

        def entry(journal, partner, amount, day, other_account):
            move = cls.env['account.move'].create({
                'move_type': 'entry', 'journal_id': journal.id, 'date': day,
                'line_ids': [
                    Command.create({'account_id': other_account.id, 'debit': amount}),
                    Command.create({'account_id': cls.receivable.id, 'credit': amount,
                                    'partner_id': partner.id}),
                ]})
            move.action_post()
            return move

        # Anita's clinic: 1000 owed before September; 500 billed, 600 received,
        # a 100 credit note in it; something billed after the period
        invoice(cls.clinic_a, 1000.0, '2026-08-10')
        invoice(cls.clinic_a, 500.0, '2026-09-05')
        entry(cls.bank, cls.clinic_a, 600.0, '2026-09-10', cls.bank.default_account_id)
        invoice(cls.clinic_a, 100.0, '2026-09-12', move_type='out_refund')
        invoice(cls.clinic_a, 999.0, '2026-10-05')
        # Binu's clinic: billed 300 in the period, nothing received
        invoice(cls.clinic_b, 300.0, '2026-09-15')
        # a customer with no salesperson
        invoice(cls.walk_in, 50.0, '2026-09-20')
        # Anita's other clinic paid everything before the period: nothing to show
        invoice(cls.settled, 200.0, '2026-08-01')
        entry(cls.bank, cls.settled, 200.0, '2026-08-02', cls.bank.default_account_id)

    def _wizard(self, **vals):
        return self.env['salesperson.summary.wizard'].create(dict({
            'period': 'custom', 'date_from': '2026-09-01', 'date_to': '2026-09-30'}, **vals))

    def _groups(self, **vals):
        return {g['name']: g for g in self._wizard(**vals)._collect()['groups']}

    def test_the_figures_per_salesperson(self):
        groups = self._groups()
        anita = groups['Anita Sales']
        self.assertEqual([anita[k] for k in ('opening', 'billed', 'received', 'adjusted', 'closing')],
                         [1000.0, 500.0, 600.0, 100.0, 800.0])
        self.assertEqual(anita['rate'], 40.0, "600 received out of 1000 + 500 to collect")
        binu = groups['Binu Sales']
        self.assertEqual((binu['billed'], binu['received'], binu['closing'], binu['rate']), (300.0, 0.0, 300.0, 0.0))
        self.assertEqual(groups['No salesperson']['closing'], 50.0)

    def test_outstanding_is_opening_plus_billed_less_what_came_in(self):
        data = self._wizard()._collect()
        for line in [data['totals']] + data['groups'] + [c for g in data['groups'] for c in g['customers']]:
            self.assertAlmostEqual(line['closing'],
                                   line['opening'] + line['billed'] - line['received'] - line['adjusted'])

    def test_ranked_by_outstanding_with_nobody_last(self):
        data = self._wizard()._collect()
        self.assertEqual([g['name'] for g in data['groups']], ['Anita Sales', 'Binu Sales', 'No salesperson'])
        self.assertEqual([g['rank'] for g in data['groups']], [1, 2, 3])
        self.assertAlmostEqual(data['groups'][0]['share'], 800.0 / 1150.0 * 100.0)
        self.assertEqual(data['totals']['closing'], 1150.0)

    def test_a_settled_customer_is_not_listed(self):
        anita = self._groups()['Anita Sales']
        self.assertEqual([c['name'] for c in anita['customers']], [self.clinic_a.name])
        self.assertEqual(anita['customers'][0]['place'], 'Manjeri')

    def test_choosing_salespersons(self):
        self.assertEqual(list(self._groups(user_ids=[Command.set(self.binu.ids)])), ['Binu Sales'])

    def test_the_dates_are_checked(self):
        with self.assertRaises(UserError):
            self._wizard(date_from='2026-10-01').action_print()

    def test_the_wizard_says_what_it_will_print(self):
        self.assertIn('3 salesperson(s)', self._wizard().summary)

    def test_the_printout(self):
        report = self.env.ref('ebshel_salesperson_summary.action_report_salesperson_summary')
        wizard = self._wizard()
        html = self.env['ir.actions.report']._render_qweb_html(report.report_name, wizard.ids)[0].decode()
        for text in ('Salesperson Summary', 'Anita Sales', 'Binu Sales', 'No salesperson', '40.0%', '800.00'):
            self.assertIn(text, html)
        self.assertNotIn(self.clinic_a.name, html, "the summary names salespersons, not customers")
        wizard.layout = 'detailed'
        html = self.env['ir.actions.report']._render_qweb_html(report.report_name, wizard.ids)[0].decode()
        self.assertIn(self.clinic_a.name, html)
        self.assertIn('Manjeri', html)
        self.assertNotIn('<tfoot', html, "wkhtmltopdf repeats a tfoot on every page: a total there "
                                         "prints under a list that goes on overleaf")

    def test_view_is_the_same_report_on_screen(self):
        action = self._wizard().action_view()
        self.assertEqual((action['report_type'], action['report_name']),
                         ('qweb-html', 'ebshel_salesperson_summary.report_salesperson_summary'))

    def test_the_workbook(self):
        if openpyxl is None:
            self.skipTest("openpyxl is not installed")
        book = openpyxl.load_workbook(io.BytesIO(self._wizard()._xlsx()))
        self.assertEqual(book.sheetnames, ['Summary', 'Customers'])
        summary = book['Summary']
        header = [cell.value for cell in summary[4]]
        self.assertEqual(header[:3], ['#', 'Salesperson', 'Customers'])
        total = [cell.value for cell in summary[summary.max_row]]
        self.assertEqual(total[1], 'Total')
        self.assertEqual(total[header.index('Outstanding')], 1150.0)
        customers = [row for row in book['Customers'].iter_rows(min_row=5, values_only=True)]
        self.assertIn(('Anita Sales', self.clinic_a.name, 'Manjeri', 1000.0, 500.0, 600.0, 100.0, 800.0, 40.0),
                      customers)
