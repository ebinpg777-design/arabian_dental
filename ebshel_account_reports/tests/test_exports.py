# -*- coding: utf-8 -*-
"""The report leaving the screen. A workbook is opened again and read the way a
spreadsheet reads it - a date must be a date, a figure a number, a percentage a
percentage - and what goes in must be what the reader chose."""
import base64
import io
from datetime import date, datetime
from unittest.mock import patch

import openpyxl

from odoo import fields
from odoo.tests import TransactionCase, tagged
from odoo.tools.misc import format_date

PROBE = 'EFR-EXPORT-PROBE'
AMOUNT = 4321.09


@tagged('post_install', '-at_install')
class TestFinReportExports(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Report = cls.env['ebshel.fin.report']
        cls.engine = cls.env['ebshel.fin.engine']
        cls.pl = cls.Report.by_key('profit_loss')
        cls.ledger = cls.Report.by_key('general_ledger')
        cls.today = fields.Date.context_today(cls.env.user)
        fy = cls.env.company.compute_fiscalyear_dates(cls.today)
        cls.year = {'date': {'preset': 'custom', 'from': fields.Date.to_string(fy['date_from']),
                             'to': fields.Date.to_string(cls.today)}}
        cls.tag = cls.env['res.partner.category'].create({'name': 'EFR export tag'})
        cls.partner = cls.env['res.partner'].create({'name': 'EFR Export Clinic', 'category_id': [(6, 0, [cls.tag.id])]})
        like = cls.env['account.account'].search([('account_type', '=', 'income')], limit=1)
        cls.income = like.copy({'name': 'EFR export income', 'code': 'EFX%s' % like.id})
        cls.invoice = cls.env['account.move'].create({
            'move_type': 'out_invoice', 'partner_id': cls.partner.id, 'invoice_date': cls.today,
            'invoice_payment_term_id': False,
            'invoice_line_ids': [(0, 0, {'name': '%s bridge' % PROBE, 'quantity': 1, 'price_unit': AMOUNT,
                                         'account_id': cls.income.id, 'tax_ids': [(5, 0, 0)]})]})
        cls.invoice.action_post()
        cls.own = dict(cls.year, partner_categories=[cls.tag.id])

    # ------------------------------------------------------------ helpers
    def _sheet(self, file, index=0):
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(file['content'])))
        return book, book.worksheets[index]

    def _head(self, sheet, first='Line'):
        for row in sheet.iter_rows():
            if row[0].value == first:
                return row[0].row, [c.value for c in row]
        self.fail("no heading row")

    def _row(self, sheet, name, start=1):
        for row in sheet.iter_rows(min_row=start):
            if row[0].value == name:
                return row
        self.fail("no row %s" % name)

    def _rev(self, options):
        data = self.pl.get_report_data(options)
        return next(l for l in data['lines'] if l.get('code') == 'REV')

    # ------------------------------------------------------------ what goes in
    def test_what_goes_in_is_the_readers_choice(self):
        rev = self._rev(self.year)
        summary = self.pl.get_export_preview(self.year, {'scope': 'summary'})
        full = self.pl.get_export_preview(self.year, {'scope': 'full'})
        kids = len(self.pl.expand_line(self.year, rev['id'])['lines'])
        screen = self.pl.get_export_preview(dict(self.year, expanded=[rev['id']]), {'scope': 'screen'})
        self.assertGreater(kids, 0)
        self.assertEqual(screen['lines'], summary['lines'] + kids, "as on the screen: the summary plus the one line unfolded")
        self.assertGreater(full['lines'], screen['lines'])
        # a summary folds what the reader had unfolded
        folded = self.pl.get_export_preview(dict(self.year, expanded=[rev['id']], unfold_all=True), {'scope': 'summary'})
        self.assertEqual(folded['lines'], summary['lines'])

    def test_an_unknown_choice_falls_back(self):
        layout = self.Report._export_layout({'scope': 'everything; DROP', 'orientation': 'diagonal', 'notes': 0})
        self.assertEqual(layout, {'scope': 'full', 'orientation': 'auto', 'filters': True, 'notes': False})

    def test_paper_turns_on_its_side_when_the_columns_need_it(self):
        self.assertFalse(self.pl.get_export_preview(self.year, {'scope': 'summary'})['landscape'])
        many = dict(self.year, comparison={'mode': 'previous', 'periods': 5})
        self.assertTrue(self.pl.get_export_preview(many, {'scope': 'summary'})['landscape'])
        self.assertFalse(self.pl.get_export_preview(many, {'scope': 'summary', 'orientation': 'portrait'})['landscape'])
        self.assertTrue(self.pl.get_export_preview(self.year, {'scope': 'summary', 'orientation': 'landscape'})['landscape'])
        ledger = dict(self.year, accounts_query=self.income.code)
        self.assertTrue(self.ledger.get_export_preview(ledger, {'scope': 'full'})['landscape'], "entries need the long side")

    # ------------------------------------------------------------ the workbook
    def test_a_figure_in_the_workbook_is_a_number(self):
        file = self.pl.export_xlsx(self.own, {'scope': 'summary'})
        book, sheet = self._sheet(file)
        at, head = self._head(sheet)
        self.assertEqual(head[0], 'Line')
        cell = self._row(sheet, 'Revenue', at)[1]
        self.assertIsInstance(cell.value, (int, float))
        self.assertAlmostEqual(cell.value, AMOUNT, 2)
        self.assertIn('#,##0.00', cell.number_format)
        self.assertEqual(file['lines'], len(self.pl.get_report_data(self.own)['lines']))

    def test_the_workbook_folds_and_prints_the_way_the_screen_reads(self):
        file = self.pl.export_xlsx(self.own, {'scope': 'full'})
        book, sheet = self._sheet(file)
        at, head = self._head(sheet)
        account = self._row(sheet, '%s %s' % (self.income.code, self.income.name), at)
        self.assertEqual(sheet.row_dimensions[account[0].row].outlineLevel, 1, "an account folds under its line")
        self.assertEqual(sheet.row_dimensions[self._row(sheet, 'Revenue', at)[0].row].outlineLevel, 0)
        self.assertFalse(sheet.sheet_properties.outlinePr.summaryBelow, "the heading of a group sits above it")
        self.assertEqual(sheet.freeze_panes, 'B%d' % (at + 1))
        self.assertEqual(sheet.print_title_rows, '$%d:$%d' % (at, at))
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(sheet.page_setup.fitToHeight, 0, "as many pages long as it takes, one page wide")

    def test_the_workbook_tells_the_filters_in_words(self):
        options = dict(self.own, journal_types=['sale'], label=PROBE)
        book, sheet = self._sheet(self.pl.export_xlsx(options, {'scope': 'summary'}))
        at, head = self._head(sheet)
        said = ' | '.join(str(sheet.cell(row=r, column=1).value) for r in range(1, at))
        self.assertIn(self.env.company.name, said)
        self.assertIn('Partner tag: EFR export tag', said)
        self.assertIn('Journal type: Sales', said)
        self.assertIn('Label contains: %s' % PROBE, said)
        self.assertIn(self.env.user.name, said)
        book, sheet = self._sheet(self.pl.export_xlsx(options, {'scope': 'summary', 'filters': False}))
        at, head = self._head(sheet)
        said = ' | '.join(str(sheet.cell(row=r, column=1).value) for r in range(1, at))
        self.assertNotIn('Partner tag', said)
        self.assertIn('Period', said, "the period is never left out")

    def test_growth_and_margins_are_real_percentages(self):
        options = dict(self.year, comparison={'mode': 'last_year', 'periods': 1}, growth=True)
        data = self.pl.get_report_data(options)
        book, sheet = self._sheet(self.pl.export_xlsx(options, {'scope': 'summary'}))
        at, head = self._head(sheet)
        self.assertEqual(head[-1], 'Change')
        checked = 0
        for line in data['lines']:
            row = self._row(sheet, line['name'], at)
            for cell, shown in zip(line['columns'], row[1:]):
                if cell['value'] is None:
                    continue
                if cell['display'] == 'percent' or (cell['display'] == 'growth' and abs(cell['value']) <= 999.9):
                    self.assertAlmostEqual(shown.value, cell['value'] / 100.0, 4)
                    self.assertIn('%', shown.number_format)
                    checked += 1
        self.assertTrue(checked, "the profit and loss has margins")

    def test_figures_in_thousands_are_divided_in_the_workbook_too(self):
        book, sheet = self._sheet(self.pl.export_xlsx(dict(self.own, unit=1000), {'scope': 'summary'}))
        at, head = self._head(sheet)
        self.assertAlmostEqual(self._row(sheet, 'Revenue', at)[1].value, AMOUNT / 1000.0, 4)
        said = ' | '.join(str(sheet.cell(row=r, column=1).value) for r in range(1, at))
        self.assertIn('Figures in: thousands', said)

    def test_a_ledger_puts_the_parts_of_an_entry_in_columns(self):
        options = dict(self.year, accounts_query=self.income.code)
        book, sheet = self._sheet(self.ledger.export_xlsx(options, {'scope': 'full'}))
        at, head = self._head(sheet)
        for wanted in ('Date', 'Entry', 'Partner', 'Label', 'Debit', 'Credit', 'Balance'):
            self.assertIn(wanted, head)
        self.assertNotIn('Account', head, "under its own account an entry does not repeat it")
        self.assertNotIn('Due', head, "a ledger line has nothing due")
        entry = next(r for r in sheet.iter_rows(min_row=at + 1) if r[head.index('Entry')].value == self.invoice.name)
        day = entry[head.index('Date')].value
        self.assertIsInstance(day, (date, datetime))
        self.assertEqual(fields.Date.to_date(day), self.today)
        self.assertEqual(entry[head.index('Partner')].value, 'EFR Export Clinic')
        self.assertIn(PROBE, entry[head.index('Label')].value)
        self.assertAlmostEqual(entry[head.index('Credit')].value, AMOUNT, 2)
        self.assertEqual(entry[0].value, '%s %s' % (self.income.code, self.income.name), "a filter on the first column keeps an account together")
        self.assertEqual(sheet.row_dimensions[entry[0].row].outlineLevel, 1)
        self.assertTrue(sheet.auto_filter.ref)

    def test_a_note_goes_with_its_line(self):
        rev = self._rev(self.own)
        self.env['ebshel.fin.report.annotation'].add_note(self.pl.id, rev['id'], 'Checked against the bank', line_name='Revenue')
        book, sheet = self._sheet(self.pl.export_xlsx(self.own, {'scope': 'summary'}))
        at, head = self._head(sheet)
        self.assertIn('Checked against the bank', self._row(sheet, 'Revenue', at)[0].comment.text)
        self.assertEqual(len(book.worksheets), 2)
        self.assertEqual(book.worksheets[1].cell(row=2, column=2).value, 'Checked against the bank')
        book, sheet = self._sheet(self.pl.export_xlsx(self.own, {'scope': 'summary', 'notes': False}))
        self.assertEqual(len(book.worksheets), 1)

    def test_an_empty_report_still_makes_a_workbook(self):
        nothing = dict(self.year, label='no entry carries this label 7e1f')
        file = self.ledger.export_xlsx(nothing, {'scope': 'full'})
        book, sheet = self._sheet(file)
        at, head = self._head(sheet)
        below = [sheet.cell(row=r, column=1).value for r in range(at + 1, sheet.max_row + 1)]
        self.assertEqual(below[0], 'Total', "a ledger of nothing still foots, to nothing")
        self.assertIn('Nothing in this period with these filters.', below)
        self.assertTrue(self.ledger.get_export_preview(nothing, {'scope': 'full'})['empty'])
        self.assertFalse(self.ledger.get_export_preview(self.year, {'scope': 'summary'})['empty'])

    # ------------------------------------------------------------ plain rows
    def test_rows_for_a_csv_are_the_same_reading(self):
        out = self.pl.export_rows(self.own, {'scope': 'summary'})
        data = self.pl.get_report_data(self.own)
        self.assertEqual(out['rows'][0][0], 'Line')
        self.assertEqual(len(out['rows']), len(data['lines']) + 1)
        self.assertEqual(len(out['levels']), len(out['rows']))
        rev = next(r for r in out['rows'] if r[0] == 'Revenue')
        self.assertAlmostEqual(rev[1], AMOUNT, 2)
        ledger = self.ledger.export_rows(dict(self.year, accounts_query=self.income.code), {'scope': 'full'})
        head = ledger['rows'][0]
        entry = next(r for r in ledger['rows'] if r[head.index('Entry')] == self.invoice.name)
        self.assertEqual(entry[head.index('Date')], format_date(self.env, self.today))

    # ------------------------------------------------------------ the items of one figure
    def test_the_items_of_a_figure_as_a_sheet(self):
        rev = self._rev(self.own)
        file = self.pl.export_items_xlsx(self.own, rev['id'], 'p0')
        self.assertEqual((file['items'], file['total']), (1, 1))
        book, sheet = self._sheet(file)
        at, head = self._head(sheet, 'Date')
        row = [c.value for c in sheet[at + 1]]
        self.assertIsInstance(row[head.index('Date')], (date, datetime))
        self.assertEqual(row[head.index('Entry')], self.invoice.name)
        self.assertEqual(row[head.index('Partner')], 'EFR Export Clinic')
        self.assertAlmostEqual(row[head.index('Credit')], AMOUNT, 2)
        self.assertAlmostEqual(row[head.index('Balance')], -AMOUNT, 2)
        self.assertEqual(row[head.index('Status')], 'Posted')
        total = sheet.cell(row=at - 1, column=head.index('Credit') + 1).value
        self.assertTrue(str(total).startswith('=SUBTOTAL(109,'), "the total follows a filter put on the sheet")
        self.assertIn('%d:' % (at + 1), str(total))
        self.assertTrue(sheet.auto_filter.ref)

    def test_the_items_sheet_follows_the_search(self):
        rev = self._rev(self.year)
        file = self.pl.export_items_xlsx(self.year, rev['id'], 'p0', search=PROBE)
        self.assertEqual(file['items'], 1)
        everything = self.pl.export_items_xlsx(self.year, rev['id'], 'p0')
        self.assertEqual(everything['total'], self.pl.get_items(self.year, rev['id'], 'p0')['total'])

    def test_the_items_sheet_is_cut_where_it_must_be_and_says_so(self):
        rev = self._rev(self.year)
        total = self.pl.get_items(self.year, rev['id'], 'p0')['total']
        if total < 2:
            second = self.invoice.copy({'invoice_date': self.today})
            second.action_post()
            total = self.pl.get_items(self.year, rev['id'], 'p0')['total']
        with patch('odoo.addons.ebshel_account_reports.models.export.ITEMS_CAP', 1):
            file = self.pl.export_items_xlsx(self.year, rev['id'], 'p0')
        self.assertEqual((file['items'], file['total']), (1, total))
        book, sheet = self._sheet(file)
        said = ' | '.join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
        self.assertIn('Cut at 1 items of %s' % total, said)

    # ------------------------------------------------------------ the PDF
    def _values(self, report, options, layout=None):
        action = report.get_pdf_action(options, layout)
        return self.env['report.ebshel_account_reports.fin_report_pdf']._get_report_values([report.id], action['data'])

    def test_the_pdf_is_cut_in_slices_the_renderer_can_carry(self):
        values = self._values(self.ledger, self.year, {'scope': 'full'})
        self.assertTrue(values['slices'])
        self.assertTrue(all(len(s) <= 250 for s in values['slices']))
        self.assertEqual(sum(len(s) for s in values['slices']), values['line_count'])
        self.assertAlmostEqual(values['name_w'] + values['col_w'] * len(values['columns']), 100.0, 0)

    def test_a_ledger_too_long_for_paper_says_so(self):
        full = self.ledger.get_export_preview(self.year, {'scope': 'full'})['lines']
        self.env['ir.config_parameter'].sudo().set_param('ebshel_account_reports.pdf_max_lines', '200')
        self.assertEqual(self.Report._pdf_cap(), 200)
        self.env['ir.config_parameter'].sudo().set_param('ebshel_account_reports.pdf_max_lines', 'many')
        self.assertEqual(self.Report._pdf_cap(), 4000, "a setting that is not a number is not obeyed")
        with patch.object(type(self.ledger), '_pdf_cap', lambda self: 3):
            values = self._values(self.ledger, self.year, {'scope': 'full'})
            preview = self.ledger.get_export_preview(self.year, {'scope': 'full'})
        self.assertGreater(full, 3)
        self.assertEqual(values['line_count'], 3)
        self.assertEqual(values['cut_total'], full)
        self.assertTrue(preview['cut'])
        self.assertEqual(preview['lines'], full, "the preview counts what the report holds, not what paper takes")

    def test_the_pdf_renders_upright_and_on_its_side(self):
        Action = self.env['ir.actions.report'].with_context(force_report_rendering=True)
        for options, layout, landscape in ((self.own, {'scope': 'summary'}, False),
                                           (dict(self.year, accounts_query=self.income.code), {'scope': 'full'}, True)):
            report = self.ledger if landscape else self.pl
            with self.subTest(landscape=landscape):
                action = report.get_pdf_action(options, layout)
                self.assertEqual(self._values(report, options, layout)['landscape'], landscape)
                pdf, kind = Action._render_qweb_pdf('ebshel_account_reports.action_fin_report_pdf', res_ids=[report.id], data=action['data'])
                self.assertTrue(pdf.startswith(b'%PDF'))

    def test_the_pdf_is_named_after_the_report(self):
        file = self.pl.with_context(force_report_rendering=True).export_pdf(self.own, {'scope': 'summary'})
        self.assertEqual(file['filename'], 'Profit and Loss - %s.pdf' % fields.Date.to_string(self.today))
        self.assertTrue(base64.b64decode(file['content']).startswith(b'%PDF'))

    def test_a_caller_from_before_the_panel_still_gets_everything(self):
        """A statement sent to a partner, a schedule: no layout in the data."""
        values = self.env['report.ebshel_account_reports.fin_report_pdf']._get_report_values(
            [self.pl.id], {'options': self.own, 'report_id': self.pl.id})
        names = [l['name'] for s in values['slices'] for l in s]
        self.assertIn('%s %s' % (self.income.code, self.income.name), names)

    # ------------------------------------------------------------ the last pass over detail rows
    def test_an_entry_is_dated_the_readers_way_and_says_nothing_twice(self):
        options = dict(self.year, accounts_query=self.income.code, expanded=['ac:%s' % self.income.id])
        lines = self.ledger.get_report_data(options)['lines']
        entry = next(l for l in lines if l.get('parts') and l['parts']['move'] == self.invoice.name)
        self.assertEqual(entry['parts']['date_text'], format_date(self.env, self.today))
        self.assertEqual(entry['parts']['date'], fields.Date.to_string(self.today), "the date itself stays a date")
        self.assertEqual(entry['parts']['account'], '', "the account is the heading above")
        self.assertEqual(entry['parts']['due'], '')
        kids = self.ledger.expand_line(dict(self.year, accounts_query=self.income.code), 'ac:%s' % self.income.id)['lines']
        entry = next(l for l in kids if l.get('parts') and l['parts']['move'] == self.invoice.name)
        self.assertEqual(entry['parts']['date_text'], format_date(self.env, self.today), "unfolding a line tidies its rows too")

    def test_under_a_partner_an_entry_does_not_repeat_the_partner(self):
        report = self.Report.by_key('partner_ledger')
        lines = report.export_rows(dict(self.year, partners=[self.partner.id]), {'scope': 'full'})
        head = lines['rows'][0]
        self.assertNotIn('Partner', head)
        self.assertIn('Entry', head)
        self.assertTrue(any(r[head.index('Entry')] == self.invoice.name for r in lines['rows']))

    def test_a_date_in_the_workbook_reads_the_readers_way(self):
        lang = self.env['res.lang']._activate_lang('en_US')
        for pattern, excel in (('%m/%d/%Y', 'mm/dd/yyyy'), ('%d/%m/%Y', 'dd/mm/yyyy'), ('%d %b %Y', 'dd mmm yyyy')):
            with self.subTest(pattern=pattern):
                lang.date_format = pattern
                self.env['res.lang'].invalidate_model()
                self.env.registry.clear_cache()
                self.assertEqual(self.Report.with_context(lang='en_US')._excel_date_format(), excel)
        options = dict(self.year, accounts_query=self.income.code)
        book, sheet = self._sheet(self.ledger.with_context(lang='en_US').export_xlsx(options, {'scope': 'full'}))
        at, head = self._head(sheet)
        entry = next(r for r in sheet.iter_rows(min_row=at + 1) if r[head.index('Entry')].value == self.invoice.name)
        self.assertEqual(entry[head.index('Date')].number_format, 'dd mmm yyyy')

    def test_the_aged_reports_give_an_eighth_of_the_first_column_to_the_figures(self):
        Pdf = self.env['report.ebshel_account_reports.fin_report_pdf']
        usual = Pdf._column_plan(8, True, False)
        for key in ('aged_receivable', 'aged_payable'):
            with self.subTest(report=key):
                report = self.Report.by_key(key)
                self.assertAlmostEqual(report.first_column, 0.875, 3)
                self.assertAlmostEqual(report.get_report_data(self.year)['report']['first_column'], 0.875, 3)
                values = Pdf._get_report_values([report.id], {'options': self.year, 'report_id': report.id, 'layout': {'scope': 'summary'}})
                count = len(values['columns'])
                plain = Pdf._column_plan(count, values['landscape'], False)
                self.assertAlmostEqual(values['name_w'], plain[0] * 0.875, 1)
                self.assertGreater(values['col_w'], plain[1], "what the names give up goes to the figures")
                self.assertAlmostEqual(values['name_w'] + values['col_w'] * count, 100.0, 0)
        self.assertEqual(self.pl.first_column, 1.0, "every other report keeps its usual width")
        self.assertEqual(usual, Pdf._column_plan(8, True, False, 1.0))

    def test_the_first_column_of_the_workbook_follows(self):
        report = self.Report.by_key('aged_receivable')
        narrow = self._sheet(report.export_xlsx(self.year, {'scope': 'summary'}))[1].column_dimensions['A'].width
        report.first_column = 1.0
        usual = self._sheet(report.export_xlsx(self.year, {'scope': 'summary'}))[1].column_dimensions['A'].width
        self.assertAlmostEqual(narrow / usual, 0.875, 1)
        report.first_column = 25          # nonsense is kept inside what a table can take
        self.assertEqual(report._first_column(), 1.6)
        report.first_column = 0
        self.assertEqual(report._first_column(), 1.0)

    def test_the_period_is_written_the_readers_way(self):
        label = self.engine.period_label(date(2026, 4, 3), date(2026, 9, 29))
        self.assertEqual(label, '%s – %s' % (format_date(self.env, date(2026, 4, 3)), format_date(self.env, date(2026, 9, 29))))
        self.assertIn(format_date(self.env, date(2026, 9, 29)), self.engine.period_label(date(2026, 4, 1), date(2026, 9, 29), True))

    # ------------------------------------------------------------ who may
    def test_every_way_out_needs_the_reader_role(self):
        nobody = self.env['res.users'].create({'name': 'Nobody', 'login': 'efr_export_nobody',
                                               'group_ids': [(6, 0, [self.env.ref('base.group_user').id])]})
        rev = self._rev(self.own)
        report = self.pl.with_user(nobody)
        for call in (lambda: report.export_xlsx(self.own), lambda: report.export_rows(self.own), lambda: report.export_pdf(self.own),
                     lambda: report.get_export_preview(self.own), lambda: report.get_pdf_action(self.own),
                     lambda: report.export_items_xlsx(self.own, rev['id'])):
            with self.assertRaises(Exception):
                call()

    def test_an_accounting_administrator_exports(self):
        admin = self.env['res.users'].create({
            'name': 'Export Admin', 'login': 'efr_export_admin',
            'group_ids': [(6, 0, [self.env.ref('base.group_user').id, self.env.ref('account.group_account_manager').id])]})
        report = self.pl.with_user(admin)
        rev = self._rev(self.own)
        self.assertEqual(report.get_export_preview(self.own, {'scope': 'summary'})['scope'], 'summary')
        self.assertTrue(report.export_xlsx(self.own, {'scope': 'full'})['content'])
        self.assertTrue(report.export_rows(self.own, {'scope': 'screen'})['rows'])
        self.assertEqual(report.export_items_xlsx(self.own, rev['id'])['items'], 1)
        self.assertEqual(report.get_pdf_action(self.own, {'scope': 'summary'})['data']['layout']['scope'], 'summary')
