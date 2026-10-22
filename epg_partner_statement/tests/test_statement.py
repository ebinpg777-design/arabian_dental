# -*- coding: utf-8 -*-
from datetime import date, datetime, timedelta

import json
import re
from urllib.parse import unquote

from odoo import fields, http
from odoo.exceptions import AccessError, UserError
from odoo.tests import tagged
from odoo.tests.common import HttpCase, TransactionCase


@tagged('post_install', '-at_install')
class TestPartnerStatement(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.partner = cls.env['res.partner'].create({
            'name': 'Statement Clinic', 'is_company': True, 'email': 'clinic@example.com',
            # No payment term, deliberately. These tests age documents by their own
            # dates, and every partner on this database now defaults to "30 Days"
            # (client, 2026-08-19) — which silently moved a 45-day-old invoice from
            # the 31-60 bucket to 1-30 and made the ageing assertions test the default
            # rather than the ageing. (client, 2026-08-27)
            'property_payment_term_id': False})
        cls.product = cls.env['product.product'].create({'name': 'Appliance', 'type': 'service', 'list_price': 100})
        cls.Engine = cls.env['epg.partner.statement']
        cls.today = fields.Date.today()

    def _invoice(self, amount, inv_date, post=True, partner=None):
        inv = self.env['account.move'].create({
            'move_type': 'out_invoice', 'partner_id': (partner or self.partner).id, 'invoice_date': inv_date,
            'invoice_date_due': inv_date, 'date': inv_date,
            'invoice_line_ids': [(0, 0, {'product_id': self.product.id, 'quantity': 1, 'price_unit': amount,
                                         'tax_ids': [(6, 0, [])]})],
        })
        if post:
            inv.action_post()
        return inv

    def _pay(self, invoice, amount, pay_date):
        return self._pay_partner(self.partner, invoice, amount, pay_date)

    def _pay_partner(self, partner, invoice, amount, pay_date):
        journal = self.env['account.journal'].search([('type', '=', 'bank'), ('company_id', '=', self.company.id)], limit=1)
        payment = self.env['account.payment'].create({
            'payment_type': 'inbound', 'partner_type': 'customer', 'partner_id': partner.id,
            'amount': amount, 'date': pay_date, 'journal_id': journal.id,
        })
        payment.action_post()
        lines = (payment.move_id.line_ids + invoice.line_ids).filtered(
            lambda l: l.account_id.account_type == 'asset_receivable' and not l.reconciled)
        lines.reconcile()
        return payment

    def _options(self, date_from, date_to, **kw):
        opts = {'statement_type': 'receivable', 'date_from': date_from, 'date_to': date_to,
                'company': self.company, 'open_items_only': False, 'show_ageing': True}
        opts.update(kw)
        return opts

    def test_ledger_opening_running_closing(self):
        old = self._invoice(100, self.today - timedelta(days=60))
        new = self._invoice(250, self.today - timedelta(days=5))
        self._pay(new, 50, self.today - timedelta(days=2))
        st = self.Engine.compute(self.partner, self._options(self.today - timedelta(days=30), self.today))[self.partner.id]
        self.assertTrue(st['has_data'])
        block = st['blocks'][0]
        self.assertEqual(block['opening'], 100.0, "invoice before the period is the opening balance")
        self.assertEqual([round(l['balance'], 2) for l in block['lines']], [350.0, 300.0])
        self.assertEqual(block['closing'], 300.0)
        self.assertEqual(block['total_debit'], 250.0)
        self.assertEqual(block['total_credit'], 50.0)

    def test_open_items_and_ageing(self):
        old = self._invoice(100, self.today - timedelta(days=45))     # 45 days overdue -> 31-60
        new = self._invoice(200, self.today - timedelta(days=5))      # 5 days -> 1-30
        paid = self._invoice(70, self.today - timedelta(days=10)); self._pay(paid, 70, self.today - timedelta(days=9))
        st = self.Engine.compute(self.partner, self._options(self.today - timedelta(days=30), self.today, open_items_only=True))[self.partner.id]
        block = st['blocks'][0]
        self.assertEqual({l['move'] for l in block['lines']}, {old.name, new.name}, "only unpaid documents, whatever their date")
        ageing = {r[0]: r[1] for r in block['ageing_rows']}
        self.assertEqual(ageing['31-60 days'], 100.0)
        self.assertEqual(ageing['1-30 days'], 200.0)
        self.assertEqual(block['ageing_total'], 300.0)
        overdue = {l['move']: l['days_overdue'] for l in block['lines']}
        self.assertEqual(overdue[old.name], 45)

    def test_payable_sign_and_type_filter(self):
        bill = self.env['account.move'].create({
            'move_type': 'in_invoice', 'partner_id': self.partner.id, 'invoice_date': self.today, 'date': self.today,
            'invoice_line_ids': [(0, 0, {'product_id': self.product.id, 'quantity': 1, 'price_unit': 40, 'tax_ids': [(6, 0, [])]})]})
        bill.action_post()
        self._invoice(100, self.today)
        recv = self.Engine.compute(self.partner, self._options(self.today, self.today))[self.partner.id]['blocks'][0]
        pay = self.Engine.compute(self.partner, self._options(self.today, self.today, statement_type='payable'))[self.partner.id]['blocks'][0]
        self.assertEqual(recv['closing'], 100.0)
        self.assertEqual(pay['closing'], 40.0, "payables are shown as a positive amount owed to the vendor")

    def test_pdf_and_xlsx_render(self):
        self._invoice(120, self.today)
        opts = self._options(self.today, self.today)
        pdf = self.Engine.render_pdf(self.partner, opts)
        # in test mode Odoo renders the HTML body instead of calling wkhtmltopdf
        self.assertTrue(pdf.startswith(b'%PDF') or self.partner.name.encode() in pdf)
        xlsx = self.Engine.render_xlsx(self.partner, opts)
        self.assertTrue(xlsx.startswith(b'PK'), "xlsx is a zip container")

    def test_wizard_all_with_balance_and_email(self):
        self._invoice(90, self.today)
        other = self.env['res.partner'].create({'name': 'Settled Clinic'})
        settled = self._invoice(10, self.today, partner=other)
        self._pay_partner(other, settled, 10, self.today)
        wiz = self.env['epg.partner.statement.wizard'].create({
            'all_with_balance': True, 'statement_type': 'receivable', 'period': 'custom',
            'date_from': self.today, 'date_to': self.today, 'company_id': self.company.id})
        self.assertIn(self.partner, wiz._partners())
        self.assertNotIn(other, wiz._partners(), "a fully paid partner is not 'with a balance'")
        mails_before = self.env['mail.mail'].search_count([])
        wiz.write({'all_with_balance': False, 'partner_ids': [(6, 0, self.partner.ids)]})
        wiz.action_send_email()
        self.assertGreater(self.env['mail.mail'].search_count([]), mails_before)
        self.assertEqual(self.partner.statement_last_sent, self.today)
        att = self.env['ir.attachment'].search([('res_model', '=', 'res.partner'), ('res_id', '=', self.partner.id), ('mimetype', '=', 'application/pdf')])
        self.assertTrue(att, "the sent statement is kept on the partner")

    def test_period_presets(self):
        d1, d2 = self.Engine.period_dates('last_month', self.company, date(2026, 3, 15))
        self.assertEqual((d1, d2), (date(2026, 2, 1), date(2026, 2, 28)))
        d1, d2 = self.Engine.period_dates('this_quarter', self.company, date(2026, 5, 10))
        self.assertEqual((d1, d2), (date(2026, 4, 1), date(2026, 6, 30)))

    def test_cron_sends_only_opted_in(self):
        self._invoice(55, self.today.replace(day=1) - timedelta(days=3))
        self.partner.write({'statement_auto_send': True})
        self.env['ir.config_parameter'].sudo().set_str('epg_partner_statement.auto_send_day', str(self.today.day))
        sent = self.Engine._cron_auto_send()
        self.assertEqual(sent, 1)
        self.assertEqual(self.Engine._cron_auto_send(), 0, "not sent twice in the same month")

    def test_the_description_names_what_the_invoice_was_for(self):
        """A clinic reconciling a statement recognises the appliance, not the sales
        order number this column used to show. (client, 2026-08-27)"""
        invoice = self._invoice(100, self.today)
        block = self.Engine.compute(
            self.partner, self._options(self.today, self.today)
        )[self.partner.id]['blocks'][0]
        line = next(l for l in block['lines'] if l['move'] == invoice.name)
        if 'product_names' in self.env['account.move']._fields:
            self.assertEqual(line['name'], invoice.product_names)
            self.assertIn(self.product.name, line['name'])

    def test_a_line_with_no_items_still_says_something(self):
        """Opening balances, payments and journal entries name no product, and the
        column must not go blank for them."""
        payment_only = self._invoice(70, self.today)
        self._pay(payment_only, 70, self.today)
        block = self.Engine.compute(
            self.partner, self._options(self.today, self.today)
        )[self.partner.id]['blocks'][0]
        for line in block['lines']:
            self.assertIsInstance(line['name'], str)

    def test_the_period_follows_the_document_date_not_the_entry_date(self):
        """A migrated opening item carries the date it had in the old system, and the
        statement always SHOWED that while FILTERING on the accounting date — so a line
        printed as 2023 was included or excluded by its 2026 migration date. 123,446 of
        this ledger's items are in exactly that position. (client, 2026-08-27)"""
        from datetime import date
        invoice = self._invoice(100, self.today)
        line = invoice.line_ids.filtered(
            lambda l: l.account_id.account_type == 'asset_receivable')
        old_day = date(2023, 10, 3)
        line.original_date = old_day
        line.flush_recordset()
        self.assertEqual(line.statement_date, old_day,
                         "the document's own date is what the statement works on")

        # It now belongs to the period it reads as...
        block = self.Engine.compute(
            self.partner, self._options(date(2023, 1, 1), date(2023, 12, 31))
        )[self.partner.id]['blocks'][0]
        self.assertIn(invoice.name, {l['move'] for l in block['lines']})

        # ...and is history by the time the entry date comes round.
        later = self.Engine.compute(
            self.partner, self._options(self.today, self.today)
        )[self.partner.id]['blocks'][0]
        self.assertNotIn(invoice.name, {l['move'] for l in later['lines']},
                         "it is opening balance now, not movement in this period")

    def test_a_line_with_no_document_date_still_uses_its_entry_date(self):
        invoice = self._invoice(50, self.today)
        line = invoice.line_ids.filtered(
            lambda l: l.account_id.account_type == 'asset_receivable')
        self.assertFalse(line.original_date)
        self.assertEqual(line.statement_date, line.date)

    def test_the_statement_still_balances_on_the_new_date_basis(self):
        """Opening + charges - payments must equal closing however the lines are dated;
        filtering on one date and cutting the opening balance on another would list a
        line in the period and count it into the opening balance as well."""
        from datetime import date, timedelta
        old = self._invoice(100, self.today - timedelta(days=200))
        old.line_ids.filtered(
            lambda l: l.account_id.account_type == 'asset_receivable'
        ).original_date = date(2023, 5, 1)
        self._invoice(200, self.today)
        block = self.Engine.compute(
            self.partner, self._options(self.today - timedelta(days=30), self.today)
        )[self.partner.id]['blocks'][0]
        self.assertAlmostEqual(
            block['opening'] + block['total_debit'] - block['total_credit'],
            block['closing'], places=2)


@tagged('post_install', '-at_install')
class TestStatementDownloadName(HttpCase):
    """The downloaded PDF must say whose statement it is.

    Odoo names a report download from `print_report_name` only when the URL
    carries the record ids; the wizard passes its partners inside `data` so the
    period travels with them, so that branch never runs and every statement
    arrived called "Statement of Account.pdf" - useless in a phone's Downloads
    folder. (client, 2026-09-01)
    """

    def setUp(self):
        super().setUp()
        self.clinic = self.env['res.partner'].create(
            {'name': 'DR NAME / EKM', 'is_company': True})

    def _download_name(self, partner_ids, date_from='2026-04-01',
                       date_to='2026-09-01'):
        """The Content-Disposition the real /report/download hands back."""
        url = ('/report/pdf/epg_partner_statement.report_statement'
               '?options=%s&context=%%7B%%7D' % json.dumps({
                   'ids': partner_ids, 'statement_type': 'receivable',
                   'date_from': date_from, 'date_to': date_to,
                   'company_id': self.env.company.id,
                   'open_items_only': True, 'show_ageing': False}))
        self.authenticate('admin', 'admin')
        response = self.url_open('/report/download', data={
            'data': json.dumps([url, 'qweb-pdf']),
            'context': '{}',
            'csrf_token': self.csrf_token(),
        })
        # RFC 5987: the header carries filename*=UTF-8''<percent-encoded>.
        raw = response.headers.get('Content-Disposition', '')
        return unquote(raw)

    def test_one_clinic_is_named_in_the_file(self):
        disposition = self._download_name([self.clinic.id])
        self.assertIn('Statement', disposition)
        # The clinic, and the period the reader chose, both survive.
        self.assertIn('DR NAME', disposition)
        self.assertIn('2026-04-01', disposition)
        self.assertNotIn('Statement of Account.pdf', disposition,
                         "the generic fallback name is the bug")

    def test_the_file_keeps_its_pdf_extension(self):
        """Without it a phone has nothing to pick a viewer from."""
        self.assertIn('.pdf', self._download_name([self.clinic.id]))

    def test_a_slash_in_the_clinic_name_cannot_become_a_path(self):
        disposition = self._download_name([self.clinic.id])
        name = disposition.split('filename', 1)[-1]
        self.assertNotIn('/', name.replace('filename*=UTF-8', ''),
                         "a clinic called 'DR X / EKM' must not invent a folder")

    def test_a_run_of_many_clinics_says_how_many(self):
        other = self.env['res.partner'].create({'name': 'Second Clinic'})
        disposition = self._download_name([self.clinic.id, other.id])
        self.assertIn('2', disposition)
        self.assertNotIn('DR NAME', disposition,
                         "a multi-clinic run belongs to no single clinic")


@tagged('post_install', '-at_install')
class TestEnteredBetweenFilter(TransactionCase):
    """Filtering on when the ENTRY was typed, not what it is dated.

    On this ledger a migrated document carries its Odoo 10 date while its
    create_date is the moment of the import, years later — so "the statement
    for August" and "only what we have entered since the import" are two
    different questions and the wizard now asks both. (client, 2026-09-02)
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.Engine = cls.env['epg.partner.statement']
        cls.partner = cls.env['res.partner'].create({
            'name': 'Entered Clinic', 'is_company': True,
            'property_payment_term_id': False})
        cls.product = cls.env['product.product'].create(
            {'name': 'Entered Appliance', 'type': 'service', 'list_price': 100})
        cls.day = date(2026, 6, 15)

    def _invoice(self, amount, created_at=None):
        """An invoice, optionally back-stamped as if it had been imported.

        create_date is not writable through the ORM, so the stamp is set in
        SQL — which is exactly what a migration does to these rows.
        """
        inv = self.env['account.move'].create({
            'move_type': 'out_invoice', 'partner_id': self.partner.id,
            'invoice_date': self.day, 'invoice_date_due': self.day,
            'date': self.day,
            'invoice_line_ids': [(0, 0, {
                'product_id': self.product.id, 'quantity': 1,
                'price_unit': amount, 'tax_ids': [(6, 0, [])]})]})
        inv.action_post()
        if created_at:
            self.env.cr.execute(
                "UPDATE account_move SET create_date = %s WHERE id = %s",
                (created_at, inv.id))
            inv.invalidate_recordset(['create_date'])
        return inv

    def _options(self, **kw):
        opts = {'statement_type': 'receivable',
                'date_from': date(2026, 1, 1), 'date_to': date(2026, 12, 31),
                'company': self.company, 'open_items_only': False,
                'show_ageing': False}
        opts.update(kw)
        return opts

    def _lines(self, **kw):
        data = self.Engine.compute(self.partner, self._options(**kw))
        blocks = data[self.partner.id]['blocks']
        return [row for block in blocks for row in block['lines']]

    # ------------------------------------------------------------ the filter
    def test_without_the_filter_both_entries_are_listed(self):
        self._invoice(100.0, created_at='2020-01-01 08:00:00')
        self._invoice(200.0)
        self.assertEqual(len(self._lines()), 2)

    def test_entered_after_leaves_the_import_out(self):
        old = self._invoice(100.0, created_at='2020-01-01 08:00:00')
        fresh = self._invoice(200.0)
        rows = self._lines(created_from=datetime(2024, 1, 1))
        names = {r['move'] for r in rows}
        self.assertIn(fresh.name, names)
        self.assertNotIn(old.name, names,
                         "an entry typed before the bound is out of scope")

    def test_entered_before_keeps_only_the_import(self):
        old = self._invoice(100.0, created_at='2020-01-01 08:00:00')
        fresh = self._invoice(200.0)
        rows = self._lines(created_to=datetime(2024, 1, 1))
        names = {r['move'] for r in rows}
        self.assertIn(old.name, names)
        self.assertNotIn(fresh.name, names)

    def test_a_window_takes_both_bounds_together(self):
        self._invoice(100.0, created_at='2020-01-01 08:00:00')
        mid = self._invoice(150.0, created_at='2022-06-01 08:00:00')
        self._invoice(200.0)
        rows = self._lines(created_from=datetime(2021, 1, 1),
                           created_to=datetime(2023, 1, 1))
        self.assertEqual({r['move'] for r in rows}, {mid.name})

    def test_the_bound_is_the_entrys_stamp_not_the_document_date(self):
        # Both invoices carry the SAME invoice_date; only the typing moment
        # differs. If the filter were reading the document date, neither the
        # include nor the exclude below could hold.
        old = self._invoice(100.0, created_at='2020-01-01 08:00:00')
        fresh = self._invoice(200.0)
        self.assertEqual(old.invoice_date, fresh.invoice_date)
        rows = self._lines(created_from=datetime(2024, 1, 1))
        self.assertEqual({r['move'] for r in rows}, {fresh.name})

    def test_an_empty_bound_means_no_bound(self):
        self._invoice(100.0, created_at='2020-01-01 08:00:00')
        self._invoice(200.0)
        self.assertEqual(len(self._lines(created_from=None, created_to=None)), 2,
                         "a falsy bound must not read as the epoch")

    # ------------------------------------------------------- through the UI
    def _wizard(self, **kw):
        return self.env['epg.partner.statement.wizard'].create(dict({
            'partner_ids': [(6, 0, self.partner.ids)],
            'date_from': date(2026, 1, 1), 'date_to': date(2026, 12, 31),
            'min_outstanding': 0.0}, **kw))

    def test_the_wizard_hands_the_bounds_to_the_engine(self):
        wiz = self._wizard(created_from=datetime(2024, 1, 1),
                           created_to=datetime(2027, 1, 1))
        options = wiz._options()
        self.assertEqual(options['created_from'], datetime(2024, 1, 1))
        self.assertEqual(options['created_to'], datetime(2027, 1, 1))

    def test_a_backwards_window_is_refused_not_silently_empty(self):
        wiz = self._wizard(created_from=datetime(2027, 1, 1),
                           created_to=datetime(2024, 1, 1))
        with self.assertRaises(UserError):
            wiz._options()

    def test_the_printed_report_carries_the_filter(self):
        self._invoice(100.0, created_at='2020-01-01 08:00:00')
        fresh = self._invoice(200.0)
        wiz = self._wizard(created_from=datetime(2024, 1, 1))
        data = wiz.action_print()['data']
        self.assertIn('created_from', data,
                      "a PDF that ignores the filter disagrees with its own screen")
        options = self.Engine.normalize_options(data)
        self.assertEqual(options['created_from'], datetime(2024, 1, 1))
        rebuilt = self.Engine.compute(self.partner, options)
        rows = [r for b in rebuilt[self.partner.id]['blocks'] for r in b['lines']]
        self.assertEqual({r['move'] for r in rows}, {fresh.name})

    def test_the_selection_skips_a_clinic_with_nothing_in_the_window(self):
        # Every entry this clinic has was typed during the import, so a run
        # bounded to "entered since 2024" must not pick it and print a page
        # with nothing on it.
        self._invoice(100.0, created_at='2020-01-01 08:00:00')
        wiz = self._wizard(created_from=datetime(2024, 1, 1),
                           min_outstanding=1.0)
        self.assertNotIn(self.partner, wiz._partners())
        wide = self._wizard(min_outstanding=1.0)
        self.assertIn(self.partner, wide._partners())

    def test_a_partner_picked_by_hand_prints_even_when_paid_up(self):
        """A vendor is usually paid up, and its statement is still its bills and
        payments: the floor drops partners from a run, never one picked by hand.
        (client, 2026-10-22: "Statement not working for vendors")"""
        inv = self._invoice(100.0)
        inv._reverse_moves([{'date': self.day, 'invoice_date': self.day}], cancel=True)
        self.assertEqual(inv.payment_state, 'reversed')
        picked = self._wizard(min_outstanding=1.0)
        self.assertIn(self.partner, picked._partners(), "picked by hand: printed")
        run = self._wizard(min_outstanding=1.0, partner_ids=[(5, 0, 0)],
                           all_with_balance=True)
        self.assertNotIn(self.partner, run._partners(),
                         "a run still leaves out what is settled")


@tagged('post_install', '-at_install')
class TestMigratedDocumentNumber(TransactionCase):
    """The Document column carries the number the clinic knows.

    Everything brought over from Odoo 10 sits on one journal entry per clinic, so
    forty rows of a statement all read "MISC/26-27/04/4601" while the doctor's own
    invoice number - the one they ask about on the phone - was pushed into
    Description. (client, 2026-09-18)
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.Engine = cls.env['epg.partner.statement']
        cls.partner = cls.env['res.partner'].create({
            'name': 'Migrated Clinic', 'is_company': True, 'property_payment_term_id': False})
        cls.product = cls.env['product.product'].create(
            {'name': 'Hawleys Appliance', 'type': 'service', 'list_price': 100})
        cls.day = date(2026, 6, 15)

    def _opening_entry(self, numbers):
        """One journal entry carrying several old documents, as the migration made it."""
        receivable = self.env['account.account'].search(
            [('account_type', '=', 'asset_receivable'), ('company_ids', 'in', self.company.id)], limit=1)
        income = self.env['account.account'].search(
            [('account_type', '=', 'income'), ('company_ids', 'in', self.company.id)], limit=1)
        lines = []
        for number, amount in numbers:
            lines.append((0, 0, {'account_id': receivable.id, 'partner_id': self.partner.id,
                                 'name': number, 'debit': amount, 'credit': 0.0,
                                 'date_maturity': self.day}))
        lines.append((0, 0, {'account_id': income.id, 'name': 'Opening balance',
                             'debit': 0.0, 'credit': sum(a for _n, a in numbers)}))
        entry = self.env['account.move'].create({
            'move_type': 'entry', 'date': self.day, 'ref': 'Opening Balance (Odoo 10 migration)',
            'line_ids': lines})
        entry.action_post()
        entry.line_ids.filtered(lambda l: l.account_id == receivable).write(
            {'original_date': self.day})
        return entry

    def _rows(self):
        options = {'statement_type': 'receivable', 'date_from': date(2026, 1, 1),
                   'date_to': date(2026, 12, 31), 'company': self.company,
                   'open_items_only': False, 'show_ageing': False}
        statement = self.Engine.compute(self.partner, options)[self.partner.id]
        return statement['blocks'][0]['lines']

    def test_the_old_invoice_number_is_the_document(self):
        entry = self._opening_entry([('OC132437', 330.0), ('CNRB/2024/3960', 660.0)])
        rows = {row['move']: row for row in self._rows()}
        self.assertEqual(set(rows), {'OC132437', 'CNRB/2024/3960'},
                         "the clinic's own numbers, not one journal entry twice")
        for row in rows.values():
            self.assertEqual(row['entry'], entry.name, "the entry is still there to be shown")
            self.assertEqual(row['name'], '', "and the number is not repeated in Description")

    def test_a_real_invoice_keeps_its_own_number_and_says_what_it_was_for(self):
        invoice = self.env['account.move'].create({
            'move_type': 'out_invoice', 'partner_id': self.partner.id,
            'invoice_date': self.day, 'invoice_date_due': self.day, 'date': self.day,
            'invoice_line_ids': [(0, 0, {'product_id': self.product.id, 'quantity': 1,
                                         'price_unit': 500.0, 'tax_ids': [(6, 0, [])]})]})
        invoice.action_post()
        [row] = [r for r in self._rows() if r['move'] == invoice.name]
        self.assertFalse(row['old_number'], "nothing migrated about it")
        self.assertIn('Hawleys Appliance', row['name'], "what the invoice was for")
        self.assertEqual(row['entry'], invoice.name)


class TestStatementOnScreen(TestPartnerStatement):
    """The statements on screen read the same numbers the PDF prints. (client, 2026-09-30)"""

    def _wizard(self, **vals):
        return self.env['epg.partner.statement.wizard'].create(dict({
            'partner_ids': [(6, 0, self.partner.ids)], 'statement_type': 'receivable', 'period': 'custom',
            'date_from': self.today, 'date_to': self.today, 'company_id': self.company.id}, **vals))

    def test_the_view_button_carries_the_wizard_filters(self):
        self._invoice(120, self.today)
        action = self._wizard(open_items_only=True).action_view()
        self.assertEqual(action['tag'], 'epg_statement_view')
        params = action['params']
        self.assertEqual(params['partner_ids'], self.partner.ids)
        self.assertEqual(params['data']['date_from'], str(self.today))
        self.assertTrue(params['data']['open_items_only'])

    def test_the_summary_and_the_statement_agree_with_the_engine(self):
        inv = self._invoice(120, self.today)
        self._pay(inv, 20, self.today)
        params = self._wizard().action_view()['params']
        summary = self.Engine.view_summary(params['partner_ids'], params['data'])
        row = summary['partners'][0]
        self.assertEqual(row['id'], self.partner.id)
        self.assertAlmostEqual(row['blocks'][0]['closing'], 100.0)
        self.assertNotIn('lines', row['blocks'][0], "the list stays light: no lines in the summary")
        self.assertAlmostEqual(summary['totals'][0]['closing'], 100.0)
        st = self.Engine.view_statement(self.partner.id, params['data'])
        lines = st['blocks'][0]['lines']
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[-1]['balance'], 100.0)
        self.assertEqual({l['move_id'] for l in lines}, {inv.id, (inv.line_ids.matched_credit_ids.credit_move_id.move_id).id})
        engine = self.Engine.compute(self.partner, self.Engine.normalize_options(params['data']))[self.partner.id]
        self.assertEqual(st['blocks'][0]['closing'], engine['blocks'][0]['closing'], "the screen and the PDF read one engine")

    def test_the_order_handed_over_is_kept(self):
        self._invoice(120, self.today)
        other = self.env['res.partner'].create({'name': 'Aardvark Clinic'})
        data = self._wizard().action_view()['params']['data']
        summary = self.Engine.view_summary([self.partner.id, other.id], data)
        self.assertEqual([p['id'] for p in summary['partners']], [self.partner.id, other.id],
                         "route order, as printed - not alphabetical")

    def test_print_and_excel_from_the_screen(self):
        self._invoice(120, self.today)
        data = self._wizard().action_view()['params']['data']
        action = self.Engine.view_print(self.partner.ids, data)
        self.assertEqual(action['report_name'], 'epg_partner_statement.report_statement')
        self.assertEqual(action['data']['date_from'], str(self.today))
        url = self.Engine.view_xlsx(self.partner.ids, data)
        self.assertTrue(url['url'].startswith('/web/content/'))

    def test_a_user_who_cannot_run_the_wizard_cannot_read_the_screen(self):
        portal = self.env['res.users'].create({
            'name': 'Statement Portal', 'login': 'statement_portal_view',
            'group_ids': [(6, 0, [self.env.ref('base.group_portal').id])]})
        self._invoice(120, self.today)
        data = self._wizard().action_view()['params']['data']
        from odoo.exceptions import AccessError
        with self.assertRaises(AccessError):
            self.Engine.with_user(portal).view_summary(self.partner.ids, data)


@tagged('post_install', '-at_install')
class TestStatementViewIsAccountingOnly(TestPartnerStatement):
    """'View statements' is for the accounting team. (client, 2026-10-06)

    Everyone internal can open this wizard -- an executive reaches it from a debtor
    row to print a clinic's statement -- so the restriction has to be on the one
    button, in the view and in the method behind it.
    """

    def _wizard(self, **vals):
        return self.env['epg.partner.statement.wizard'].create(dict({
            'partner_ids': [(6, 0, self.partner.ids)], 'statement_type': 'receivable', 'period': 'custom',
            'date_from': self.today, 'date_to': self.today, 'company_id': self.company.id}, **vals))

    def _user(self, login, *group_xmlids):
        return self.env['res.users'].create({
            'name': login, 'login': login,
            'group_ids': [(6, 0, [self.env.ref(g).id for g in
                                  ('base.group_user',) + group_xmlids])]})

    def _footer_buttons(self, user):
        """The button names the given user is actually served in the wizard form."""
        arch = self.env['epg.partner.statement.wizard'].with_user(user).get_view(
            self.env.ref('epg_partner_statement.view_partner_statement_wizard').id, 'form')['arch']
        return set(re.findall(r'<button[^>]*name="(\w+)"', arch))

    # ------------------------------------------------------------------ the view
    def test_the_button_is_not_served_to_a_user_outside_accounting(self):
        buttons = self._footer_buttons(self._user('stmt_plain'))
        self.assertNotIn('action_view', buttons,
                         "'View statements' must not reach a user with no accounting rights")
        self.assertIn('action_print', buttons,
                      "printing a statement is not what is being restricted")

    def test_the_button_is_served_to_invoicing_and_to_readonly(self):
        for login, group in (('stmt_inv', 'account.group_account_invoice'),
                             ('stmt_ro', 'account.group_account_readonly')):
            with self.subTest(group=group):
                self.assertIn('action_view', self._footer_buttons(self._user(login, group)),
                              "the accounting team must still get the button")

    # ------------------------------------------------------------------ the method
    def test_calling_it_anyway_is_refused(self):
        """Hiding a button does not stop an RPC call.

        The message is asserted, not just the exception type: a user outside
        accounting cannot read account.move.line either, so this raises AccessError
        with or without the guard, and a bare assertRaises passes vacuously.
        """
        self._invoice(120, self.today)
        wizard = self._wizard()
        plain = self._user('stmt_plain_rpc')
        with self.assertRaises(AccessError) as caught:
            wizard.with_user(plain).action_view()
        self.assertIn('accounting team', str(caught.exception),
                      "must be refused by the guard, before it ever reads the ledger")

    def test_the_accounting_team_can_call_it(self):
        self._invoice(120, self.today)
        for login, group in (('stmt_inv_rpc', 'account.group_account_invoice'),
                             ('stmt_ro_rpc', 'account.group_account_readonly')):
            with self.subTest(group=group):
                action = self._wizard().with_user(self._user(login, group)).action_view()
                self.assertEqual(action['tag'], 'epg_statement_view')

    def test_the_guard_is_on_this_button_only(self):
        """Print is untouched: for a user outside accounting it fails where it always
        did, on reading account.move.line, and not with the new refusal.

        (77 of the active internal users on the 2026-08-28 copy hold no accounting
        group at all, so the wizard was already unusable for them -- this change only
        stops offering them the button.)
        """
        self._invoice(120, self.today)
        plain = self._user('stmt_plain_print')
        with self.assertRaises(AccessError) as caught:
            self._wizard().with_user(plain).action_print()
        self.assertNotIn('accounting team', str(caught.exception),
                         "print must not start refusing people for the new reason")

    def test_the_accounting_team_can_still_print(self):
        self._invoice(120, self.today)
        inv = self._user('stmt_inv_print', 'account.group_account_invoice')
        self.assertTrue(self._wizard().with_user(inv).action_print())


@tagged('post_install', '-at_install')
class TestItemwiseStatement(TransactionCase):
    """The Odoo 17 "Itemwise Account Statement Report" as an option of the statement:
    each invoice's items under its row, on a landscape page. (client, 2026-10-21)"""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.partner = cls.env['res.partner'].create({
            'name': 'Itemwise Clinic', 'is_company': True, 'property_payment_term_id': False})
        cls.crown = cls.env['product.product'].create({'name': 'Metal Ceramic Crown', 'type': 'service'})
        cls.model = cls.env['product.product'].create({'name': 'Printed Model', 'type': 'service'})
        cls.Engine = cls.env['epg.partner.statement']
        cls.today = fields.Date.today()
        cls.has_patient = 'patient' in cls.env['account.move.line']._fields

    def _move(self, lines, move_type='out_invoice', **vals):
        move = self.env['account.move'].create(dict({
            'move_type': move_type, 'partner_id': self.partner.id, 'invoice_date': self.today,
            'date': self.today,
            'invoice_line_ids': [(0, 0, dict({'tax_ids': [(6, 0, [])]}, **line)) for line in lines],
        }, **vals))
        move.action_post()
        return move

    def _line(self, product, qty, price, patient=None):
        line = {'product_id': product.id, 'quantity': qty, 'price_unit': price}
        if patient and self.has_patient:
            line['patient'] = patient
        return line

    def _options(self, **kw):
        return dict({'statement_type': 'receivable', 'date_from': self.today, 'date_to': self.today,
                     'company': self.company, 'open_items_only': False, 'show_ageing': False,
                     'itemwise': True}, **kw)

    def _block(self, **kw):
        return self.Engine.compute(self.partner, self._options(**kw))[self.partner.id]['blocks'][0]

    def _wizard(self, **vals):
        return self.env['epg.partner.statement.wizard'].create(dict({
            'partner_ids': [(6, 0, self.partner.ids)], 'statement_type': 'receivable', 'period': 'custom',
            'date_from': self.today, 'date_to': self.today, 'company_id': self.company.id}, **vals))

    def test_the_items_are_listed_under_their_invoice(self):
        invoice = self._move([self._line(self.crown, 2, 650, 'Latha'), self._line(self.model, 1, 50, 'Latha')])
        row = next(r for r in self._block()['lines'] if r['entry'] == invoice.name)
        self.assertEqual([(i['name'], i['quantity'], i['price_unit'], i['amount']) for i in row['items']],
                         [('Metal Ceramic Crown', 2.0, 650.0, 1300.0), ('Printed Model', 1.0, 50.0, 50.0)])
        self.assertEqual(sum(i['amount'] for i in row['items']), row['debit'],
                         "the items add up to the row: the amount is with tax")
        self.assertEqual(row['name'], '', "the items are the description now, not repeated above them")
        if self.has_patient:
            self.assertEqual({i['patient'] for i in row['items']}, {'Latha'})

    def test_a_line_that_names_nothing_is_not_an_item(self):
        invoice = self._move([self._line(self.crown, 1, 650),
                              {'name': '/', 'quantity': 1, 'price_unit': 0},
                              {'name': 'Courier charge', 'quantity': 1, 'price_unit': 40}])
        row = next(r for r in self._block()['lines'] if r['entry'] == invoice.name)
        self.assertEqual([i['name'] for i in row['items']], ['Metal Ceramic Crown', 'Courier charge'])

    def test_the_quantity_total_nets_off_credit_notes(self):
        self._move([self._line(self.crown, 3, 650)])
        self._move([self._line(self.crown, 1, 650)], move_type='out_refund')
        self.assertEqual(self._block()['total_qty'], 2.0)

    def test_an_invoice_on_two_due_dates_lists_its_items_once(self):
        term = self.env['account.payment.term'].create({'name': 'Two halves', 'line_ids': [
            (0, 0, {'value': 'percent', 'value_amount': 50.0, 'nb_days': 0}),
            (0, 0, {'value': 'percent', 'value_amount': 50.0, 'nb_days': 30})]})
        invoice = self._move([self._line(self.crown, 1, 600)], invoice_payment_term_id=term.id)
        rows = [r for r in self._block()['lines'] if r['entry'] == invoice.name]
        self.assertEqual(len(rows), 2)
        self.assertEqual(sum(len(r['items']) for r in rows), 1)

    def test_the_standard_statement_is_unchanged(self):
        invoice = self._move([self._line(self.crown, 2, 650)])
        row = next(r for r in self._block(itemwise=False)['lines'] if r['entry'] == invoice.name)
        self.assertNotIn('items', row)
        action = self._wizard().action_print()
        self.assertEqual(action['report_name'], 'epg_partner_statement.report_statement')
        self.assertNotIn('itemwise', action['data'])

    def test_the_wizard_prints_the_landscape_report(self):
        self._move([self._line(self.crown, 2, 650)])
        action = self._wizard(itemwise=True).action_print()
        self.assertEqual(action['report_name'], 'epg_partner_statement.report_statement_itemwise')
        self.assertTrue(action['data']['itemwise'])
        report = self.env.ref('epg_partner_statement.action_report_partner_statement_itemwise')
        self.assertEqual(report.paperformat_id.orientation, 'Landscape')

    def test_the_printed_page_carries_the_items_and_the_columns(self):
        self._move([self._line(self.crown, 2, 650, 'Latha')])
        report = self.env.ref('epg_partner_statement.action_report_partner_statement_itemwise')
        data = self.Engine.report_data(self.partner, self._options())
        html = self.env['ir.actions.report']._render_qweb_html(report.report_name, self.partner.ids, data=data)[0]
        html = html.decode()
        for text in ('Item-wise', 'Metal Ceramic Crown', '>Qty<', '>Rate<', '>Amount<', '1,300.00'):
            self.assertIn(text, html)
        if self.has_patient:
            self.assertIn('Latha', html)
        # the standard print is the standard print
        plain = self.env['ir.actions.report']._render_qweb_html(
            'epg_partner_statement.report_statement', self.partner.ids,
            data=self.Engine.report_data(self.partner, self._options(itemwise=False)))[0].decode()
        self.assertNotIn('>Qty<', plain)
        self.assertNotIn('Item-wise', plain)

    def test_the_excel_has_the_item_columns_and_rows(self):
        import io
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        self._move([self._line(self.crown, 2, 650), self._line(self.model, 1, 50)])
        sheet = openpyxl.load_workbook(io.BytesIO(self.Engine.render_xlsx(self.partner, self._options()))).active
        rows = [[c for c in r] for r in sheet.iter_rows(values_only=True)]
        header = next(r for r in rows if r[0] == 'Date')
        for title in ('Qty', 'Rate', 'Amount', 'Balance'):
            self.assertIn(title, header)
        qty, amount, desc = header.index('Qty'), header.index('Amount'), header.index('Description')
        items = [(r[desc], r[qty], r[amount]) for r in rows if r[qty] and r[desc] and r[0] is None]
        self.assertIn(('Metal Ceramic Crown', 2, 1300), items)
        self.assertIn(('Printed Model', 1, 50), items)
        name = self.Engine.xlsx_name(self.partner, self._options())
        self.assertTrue(name.startswith('Item-wise Statement'))

    def test_the_screen_shows_the_items_and_prints_item_wise(self):
        self._move([self._line(self.crown, 2, 650)])
        data = self._wizard(itemwise=True).action_view()['params']['data']
        self.assertTrue(data['itemwise'])
        st = self.Engine.view_statement(self.partner.id, data)
        line = st['blocks'][0]['lines'][0]
        self.assertEqual(line['items'][0]['name'], 'Metal Ceramic Crown')
        self.assertEqual(st['blocks'][0]['total_qty'], 2.0)
        self.assertTrue(self.Engine.view_summary(self.partner.ids, data)['itemwise'])
        action = self.Engine.view_print(self.partner.ids, data)
        self.assertEqual(action['report_name'], 'epg_partner_statement.report_statement_itemwise')

    # ------------------------------------------------------------ patient-wise
    # Designed new (client, 2026-10-21): the Odoo 17 server had no patient-wise statement.

    def _patient_setup(self):
        if not self.has_patient:
            self.skipTest("no patient on invoice lines in this database")
        self._move([self._line(self.crown, 2, 650, 'Latha'), self._line(self.model, 1, 50, 'Latha')])
        self._move([self._line(self.crown, 1, 650, 'latha '), self._line(self.crown, 1, 650, 'Ramesh'),
                    self._line(self.model, 1, 50)])
        self._move([self._line(self.model, 1, 50, 'Latha')], move_type='out_refund')
        journal = self.env['account.journal'].search(
            [('type', '=', 'bank'), ('company_id', '=', self.company.id)], limit=1)
        payment = self.env['account.payment'].create({
            'payment_type': 'inbound', 'partner_type': 'customer', 'partner_id': self.partner.id,
            'amount': 500.0, 'date': self.today, 'journal_id': journal.id})
        payment.action_post()

    def test_patientwise_groups_the_works_by_patient(self):
        self._patient_setup()
        block = self._block(itemwise=False, patientwise=True)
        names = [p['name'] for p in block['patients']]
        self.assertEqual(names, ['Latha', 'Ramesh', 'Without patient name'],
                         "alphabetical, one Latha however it was typed, the nameless last")
        latha = block['patients'][0]
        self.assertEqual(latha['total'], 1300.0 + 50.0 + 650.0 - 50.0, "a credit note is a negative work")
        self.assertEqual((latha['documents'], latha['qty']), (3, 3.0))
        self.assertAlmostEqual(block['patients_total'], 1950.0 + 650.0 + 50.0)
        self.assertTrue(any(r['credit'] == 500.0 for r in block['other_rows']),
                        "a payment belongs to no patient: it is listed apart")

    def test_one_patients_statement(self):
        self._patient_setup()
        block = self._block(itemwise=False, patientwise=True, patient='rame')
        self.assertEqual([p['name'] for p in block['patients']], ['Ramesh'])

    def test_the_patientwise_print(self):
        self._patient_setup()
        action = self._wizard(patientwise=True).action_print()
        self.assertEqual(action['report_name'], 'epg_partner_statement.report_statement',
                         "the portrait statement: six columns fit")
        self.assertTrue(action['data']['patientwise'])
        render = self.env['ir.actions.report']._render_qweb_html
        html = render(action['report_name'], self.partner.ids, data=action['data'])[0].decode()
        for text in ('Patient-wise', 'Works by patient', 'Latha', 'Ramesh', 'Payments, credits and other entries',
                     'epg-summary'):
            self.assertIn(text, html)
        data = self._wizard(patientwise=True, patient='Ramesh').action_print()['data']
        self.assertEqual(data['patient'], 'Ramesh')
        html = render(action['report_name'], self.partner.ids, data=data)[0].decode()
        self.assertIn('Ramesh', html)
        self.assertNotIn('Latha', html)
        self.assertNotIn('class="epg-summary"', html, "the clinic's balance is not one patient's")

    def test_the_patientwise_excel(self):
        import io
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        self._patient_setup()
        options = self._options(itemwise=False, patientwise=True)
        sheet = openpyxl.load_workbook(io.BytesIO(self.Engine.render_xlsx(self.partner, options))).active
        rows = [r for r in sheet.iter_rows(values_only=True)]
        self.assertTrue(any(r[0] == 'Patient' and r[3] == 'Item' for r in rows))
        self.assertTrue(any(r[0] == 'All patients' and r[6] == 2650.0 for r in rows))
        self.assertTrue(self.Engine.xlsx_name(self.partner, options).startswith('Patient-wise Statement'))

    def test_item_wise_and_patient_wise_are_one_or_the_other(self):
        from odoo.tests import Form
        with Form(self.env['epg.partner.statement.wizard']) as wizard:
            wizard.itemwise = True
            wizard.patientwise = True
            self.assertFalse(wizard.itemwise)
            wizard.itemwise = True
            self.assertFalse(wizard.patientwise)
